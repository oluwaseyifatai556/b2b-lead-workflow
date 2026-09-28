"""Write the human-facing outputs: summary.md and a formatted Excel workbook."""

from __future__ import annotations

import math
from collections import Counter
from pathlib import Path
from typing import TYPE_CHECKING

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from leadflow.clean import is_missing
from leadflow.reference import FIELD_LABELS

if TYPE_CHECKING:
    from leadflow.config import IcpConfig
    from leadflow.pipeline import RunResult

# Friendly Excel headers for the snake_case CSV columns.
HEADERS = {
    **FIELD_LABELS,
    "rank": "Rank",
    "tier": "Tier",
    "score": "Score",
    "score_reasons": "Why this score",
    "missing_fields": "Missing (scored) fields",
    "needs_enrichment": "Needs enrichment?",
    "reject_reason": "Why rejected",
    "company_name_source": "Company name source",
    "profile_completeness": "Profile complete %",
    "source_rows": "Original row(s)",
    "duplicates_merged": "Duplicates merged",
}
COLUMN_WIDTHS = {
    "score_reasons": 60, "reject_reason": 45, "email": 32, "contact_title": 28, "company_name": 28,
    "missing_fields": 30,
}
HEADER_FILL = PatternFill("solid", fgColor="1F3864")
HEADER_FONT = Font(bold=True, color="FFFFFF")
SECTION_FONT = Font(bold=True, size=12, color="1F3864")


def reason_counts(rejected: pd.DataFrame) -> list[tuple[str, int]]:
    counter: Counter[str] = Counter()
    for reasons in rejected.get("reject_reason", []):
        counter.update(r for r in str(reasons).split("; ") if r)
    return sorted(counter.items(), key=lambda kv: (-kv[1], kv[0]))


# Must describe score.rank_leads exactly. The README quotes it and a test checks they match.
TIE_BREAK_TEXT = (
    "Leads are sorted by score (highest first). Ties are broken by profile completeness "
    "(more complete first), then date added (newest first), then original row order."
)


def tier_ranges(config: IcpConfig) -> list[tuple[str, int, int]]:
    """(tier, lowest score, highest score) for each tier, best tier first. Scores are whole numbers."""
    ranges, high = [], 100
    for tier in config.tiers:  # sorted highest threshold first
        low = math.ceil(tier.min_score)
        ranges.append((tier.name, low, high))
        high = low - 1
    return ranges


def tier_counts(result: RunResult) -> list[tuple[str, str, int]]:
    counts = result.ranked["tier"].value_counts()
    return [(t.name, t.description, int(counts.get(t.name, 0))) for t in result.config.tiers]


def missing_data_stats(result: RunResult) -> dict:
    ranked = result.ranked
    field_counts: Counter[str] = Counter()
    for fields in ranked["missing_fields"]:
        field_counts.update(f for f in str(fields).split("; ") if f)
    flagged = ranked[ranked["needs_enrichment"]]
    return {
        "with_gaps": int((ranked["missing_fields"] != "").sum()),
        "needs_enrichment": len(flagged),
        "by_tier": [(t.name, int((flagged["tier"] == t.name).sum())) for t in result.config.tiers],
        "fields": sorted(field_counts.items(), key=lambda kv: (-kv[1], kv[0])),
    }


MISSING_DATA_NOTE = (
    "A blank field earns no points, so a low score can mean *we don't know* rather than *poor fit*. "
    "`needs_enrichment` is Yes when filling the lead's blank fields could move it up a tier."
)


def _pct(part: int, whole: int) -> str:
    return f"{100 * part / whole:.0f}%" if whole else "0%"


def _md_cell(value: object) -> str:
    return "" if is_missing(value) else str(value).replace("|", "/")


def build_summary_markdown(result: RunResult, top_n: int = 10) -> str:
    s = result.stats
    lines = [
        "# Lead list summary",
        "",
        f"**Client ICP:** {result.config.client_name}  ",
        f"**Input file:** `{result.input_name}`  ",
        f"**Generated:** {result.run_date.isoformat()}",
        "",
        "## At a glance",
        "",
        "| | Leads |",
        "|---|---:|",
        f"| Rows in the original file | {s['rows_in']} |",
        f"| Duplicate rows merged away | {s['duplicates_removed']} |",
        f"| Unique contacts | {s['unique_contacts']} |",
        f"| Rejected (see reasons below) | {s['rejected']} |",
        f"| **Ranked and ready to use** | **{s['ranked']}** |",
        "",
        "## Tier breakdown",
        "",
        "| Tier | Score range | What it means | Leads | Share |",
        "|---|---|---|---:|---:|",
    ]
    ranges = {name: f"{low}-{high}" for name, low, high in tier_ranges(result.config)}
    for name, description, count in tier_counts(result):
        lines.append(
            f"| {name} | {ranges[name]} | {_md_cell(description)} | {count} | {_pct(count, s['ranked'])} |"
        )
    lines += ["", f"Cut-offs come from icp.yaml. {TIE_BREAK_TEXT}"]

    m = missing_data_stats(result)
    lines += [
        "",
        "## Leads with missing data",
        "",
        MISSING_DATA_NOTE,
        "",
        f"- **{m['with_gaps']}** of {s['ranked']} ranked leads are blank in at least one scored field "
        "(see the `missing_fields` column).",
        f"- **{m['needs_enrichment']}** could move up a tier if those gaps were filled: worth researching "
        "before writing them off.",
        "",
        "| Tier now | Could move up |",
        "|---|---:|",
    ]
    lines += [f"| {name} | {count} |" for name, count in m["by_tier"]]
    if m["fields"]:
        lines += ["", "| Most often missing | Leads |", "|---|---:|"]
        lines += [f"| {field} | {count} |" for field, count in m["fields"]]

    lines += ["", f"## Top {top_n} leads", ""]
    top = result.ranked.head(top_n)
    if top.empty:
        lines.append("_No leads passed the checks._")
    else:
        lines += [
            "| Rank | Tier | Score | Contact | Job title | Company | Country |",
            "|---:|---|---:|---|---|---|---|",
        ]
        for r in top.to_dict("records"):
            lines.append(
                f"| {r['rank']} | {r['tier']} | {r['score']} | {_md_cell(r['full_name'])} | "
                f"{_md_cell(r['contact_title'])} | {_md_cell(r['company_name'])} | {_md_cell(r['country'])} |"
            )

    lines += ["", "## Why leads were rejected", ""]
    reasons = reason_counts(result.rejected)
    if reasons:
        lines += ["| Reason | Leads |", "|---|---:|"]
        lines += [f"| {_md_cell(reason)} | {count} |" for reason, count in reasons]
        lines += ["", "_A lead can have more than one reason._"]
    else:
        lines.append("_No leads were rejected._")

    lines += [
        "",
        "## Data quality fixes",
        "",
        "How many values in each column were blank, tidied into a standard format, or unreadable "
        "(unreadable values are left blank; invalid emails are kept so you can see them).",
        "",
        "| Column | Blank | Standardized | Unreadable |",
        "|---|---:|---:|---:|",
    ]
    for q in result.quality.to_dict("records"):
        label = HEADERS.get(q["field"], q["field"])
        lines.append(f"| {label} | {q['blank']} | {q['standardized']} | {q['unreadable']} |")

    lines += ["", "## Scoring rules used (from icp.yaml)", "", "| Rule | Points |", "|---|---:|"]
    lines += [f"| {_md_cell(rule.name)} | {rule.points:+d} |" for rule in result.config.rules]
    lines += [
        "",
        f"Scores are points earned divided by the maximum possible ({result.config.max_points}), "
        "shown out of 100.",
        "",
    ]
    return "\n".join(lines)


def write_excel(result: RunResult, path: Path) -> None:
    wb = Workbook()
    ranked_ws = wb.active
    ranked_ws.title = "Ranked"
    tier_colors = {t.name: t.color for t in result.config.tiers}
    _write_table(ranked_ws, result.ranked, tier_colors)
    _write_table(wb.create_sheet("Rejected"), result.rejected, tier_colors)
    _write_summary_sheet(wb.create_sheet("Summary"), result)
    wb.save(path)


def _write_table(ws: Worksheet, df: pd.DataFrame, tier_colors: dict[str, str]) -> None:
    columns = list(df.columns)
    ws.append([HEADERS.get(c, c) for c in columns])
    for record in df.to_dict("records"):
        ws.append([_excel_value(record[c]) for c in columns])

    for cell in ws[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    ws.freeze_panes = "A2"
    if columns:
        ws.auto_filter.ref = f"A1:{get_column_letter(len(columns))}{max(ws.max_row, 1)}"

    for idx, col in enumerate(columns, start=1):
        letter = get_column_letter(idx)
        ws.column_dimensions[letter].width = COLUMN_WIDTHS.get(col, max(12, min(24, len(HEADERS.get(col, col)) + 4)))
        fmt = {"employee_count": "#,##0", "annual_revenue": "#,##0", "created_date": "yyyy-mm-dd"}.get(col)
        for row in range(2, ws.max_row + 1):
            cell = ws.cell(row=row, column=idx)
            if fmt:
                cell.number_format = fmt
            if col == "tier" and cell.value in tier_colors:
                cell.fill = PatternFill("solid", fgColor=tier_colors[cell.value])
                cell.font = Font(bold=True)
                cell.alignment = Alignment(horizontal="center")


def _excel_value(value: object) -> object:
    if is_missing(value):
        return None
    if isinstance(value, bool):
        return "Yes" if value else "No"
    return value


def _write_summary_sheet(ws: Worksheet, result: RunResult) -> None:
    s = result.stats
    ws.column_dimensions["A"].width = 44
    ws.column_dimensions["B"].width = 44
    ws.column_dimensions["C"].width = 14
    ws.column_dimensions["D"].width = 14
    ws.column_dimensions["E"].width = 14

    def section(title: str) -> None:
        if ws.max_row > 1:
            ws.append([])
        ws.append([title])
        ws.cell(row=ws.max_row, column=1).font = SECTION_FONT

    def header(*labels: str) -> None:
        ws.append(list(labels))
        for cell in ws[ws.max_row]:
            cell.fill = HEADER_FILL
            cell.font = HEADER_FONT

    ws.append(["Lead list summary"])
    ws["A1"].font = Font(bold=True, size=16, color="1F3864")
    ws.append(["Client ICP", result.config.client_name])
    ws.append(["Input file", result.input_name])
    ws.append(["Generated", result.run_date.isoformat()])

    section("At a glance")
    header("Step", "Leads")
    for label, key in [
        ("Rows in the original file", "rows_in"),
        ("Duplicate rows merged away", "duplicates_removed"),
        ("Unique contacts", "unique_contacts"),
        ("Rejected", "rejected"),
        ("Ranked and ready to use", "ranked"),
    ]:
        ws.append([label, s[key]])
    ws.cell(row=ws.max_row, column=1).font = Font(bold=True)
    ws.cell(row=ws.max_row, column=2).font = Font(bold=True)

    section("Tier breakdown")
    header("Tier", "What it means", "Leads", "Share", "Score range")
    colors = {t.name: t.color for t in result.config.tiers}
    ranges = {name: f"{low}-{high}" for name, low, high in tier_ranges(result.config)}
    for name, description, count in tier_counts(result):
        ws.append([name, description, count, _pct(count, s["ranked"]), ranges[name]])
        ws.cell(row=ws.max_row, column=1).fill = PatternFill("solid", fgColor=colors[name])
        ws.cell(row=ws.max_row, column=1).font = Font(bold=True)
    ws.append([TIE_BREAK_TEXT])

    section("Leads with missing data")
    ws.append([MISSING_DATA_NOTE.replace("*", "").replace("`", "")])
    m = missing_data_stats(result)
    header("", "Leads")
    ws.append(["Blank in at least one scored field", m["with_gaps"]])
    ws.append(["Could move up a tier if filled (needs enrichment)", m["needs_enrichment"]])
    header("Tier now", "Could move up")
    for name, count in m["by_tier"]:
        ws.append([name, count])
    header("Most often missing", "Leads")
    for field, count in m["fields"]:
        ws.append([field, count])

    section("Why leads were rejected")
    header("Reason", "Leads")
    for reason, count in reason_counts(result.rejected) or [("No leads were rejected", 0)]:
        ws.append([reason, count])

    section("Data quality fixes")
    header("Column", "Blank", "Standardized", "Unreadable")
    for q in result.quality.to_dict("records"):
        ws.append([HEADERS.get(q["field"], q["field"]), q["blank"], q["standardized"], q["unreadable"]])

    section("Scoring rules used (from icp.yaml)")
    header("Rule", "Points")
    for rule in result.config.rules:
        ws.append([rule.name, rule.points])
    for row in ws.iter_rows(min_col=2, max_col=4):
        for cell in row:
            if isinstance(cell.value, (int, float)):
                cell.alignment = Alignment(horizontal="right")
