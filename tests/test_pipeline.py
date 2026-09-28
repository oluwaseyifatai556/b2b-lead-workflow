from datetime import date
from pathlib import Path

import pandas as pd
import pytest
from openpyxl import load_workbook

from leadflow.__main__ import main
from leadflow.clean import InputError
from leadflow.pipeline import RANKED_COLUMNS, REJECTED_COLUMNS, process, read_leads_csv, run

ROOT = Path(__file__).resolve().parents[1]
RUN_DATE = date(2026, 1, 31)


def test_process_end_to_end(messy_leads, config):
    result = process(messy_leads, config, run_date=RUN_DATE)

    assert result.stats == {"rows_in": 7, "duplicates_removed": 2, "unique_contacts": 5, "rejected": 1, "ranked": 4}
    assert list(result.ranked.columns) == RANKED_COLUMNS
    assert list(result.rejected.columns) == REJECTED_COLUMNS
    assert {"missing_fields", "needs_enrichment"} <= set(RANKED_COLUMNS)

    top = result.ranked.iloc[0]
    assert top["full_name"] == "Jane Doe"
    assert top["score"] == 100 and top["tier"] == "A"
    assert top["source_rows"] == "2; 3; 4"
    assert top["city"] == "San Francisco"   # filled in from a duplicate
    assert top["lead_source"] == "LinkedIn"
    assert top["email"] == "jane.doe@acme.example"

    assert result.ranked["full_name"].tolist() == ["Jane Doe", "Raj Patel", "Bo Chen", "Al Ng"]
    # Bo had no company name but a company email, so the name is inferred instead of rejecting.
    bo = result.ranked.set_index("full_name").loc["Bo Chen"]
    assert (bo["company_name"], bo["company_name_source"]) == ("Initech", "inferred")
    assert result.rejected["reject_reason"].tolist() == ["Invalid email"]


def test_summary_prints_tier_cutoffs_and_tie_breaks(messy_leads, config):
    from leadflow.report import build_summary_markdown

    md = build_summary_markdown(process(messy_leads, config, run_date=RUN_DATE))
    # Fixture tiers: A >= 80, B >= 50, C >= 0
    assert "| Tier | Score range | What it means | Leads | Share |" in md
    assert "| A | 80-100 |" in md and "| B | 50-79 |" in md and "| C | 0-49 |" in md
    assert (
        "Leads are sorted by score (highest first). Ties are broken by profile completeness "
        "(more complete first), then date added (newest first), then original row order."
    ) in md


def test_readme_matches_shipped_tier_cutoffs():
    from leadflow.config import load_config
    from leadflow.report import TIE_BREAK_TEXT, tier_ranges

    readme = (ROOT / "README.md").read_text(encoding="utf-8")
    for name, low, high in tier_ranges(load_config(ROOT / "icp.yaml")):
        assert f"| **{name}** | {low}-{high} |" in readme, f"README tier table is out of date for tier {name}"
    assert TIE_BREAK_TEXT in readme


def test_unreadable_phone_scores_higher_than_blank_phone_end_to_end():
    from leadflow.config import load_config
    from tests.conftest import raw_frame

    lead = {"company_name": "Acme", "website": "acme.example", "contact_last_name": "Doe", "country": "US",
            "industry": "SaaS", "employee_count": "120", "contact_title": "VP Sales"}
    raw = raw_frame([
        {**lead, "contact_first_name": "Blank", "email": "blank@acme.example", "phone": ""},
        {**lead, "contact_first_name": "Junk", "email": "junk@acme.example", "phone": "+1-525-8369"},
    ])
    ranked = process(raw, load_config(ROOT / "icp.yaml"), run_date=RUN_DATE).ranked.set_index("contact_first_name")
    assert "No phone number" in ranked.loc["Blank", "score_reasons"]
    assert "No phone number" not in ranked.loc["Junk", "score_reasons"]
    assert ranked.loc["Junk", "score"] - ranked.loc["Blank", "score"] == 5
    assert ranked.loc["Junk", "missing_fields"].endswith("Phone (unreadable)")


def test_run_writes_all_outputs(tmp_path, messy_leads):
    csv_path = tmp_path / "leads.csv"
    messy_leads.to_csv(csv_path, index=False)
    out = tmp_path / "out"

    run(csv_path, ROOT / "icp.yaml", out, run_date=RUN_DATE)

    for name in ("ranked_leads.csv", "rejected_leads.csv", "summary.md", "ranked_leads.xlsx"):
        assert (out / name).exists(), name

    summary = (out / "summary.md").read_text(encoding="utf-8")
    assert "Generated:** 2026-01-31" in summary
    assert "| Rows in the original file | 7 |" in summary
    assert "Email is missing or not a valid address" in summary
    assert "## Leads with missing data" in summary
    assert "could move up a tier" in summary

    wb = load_workbook(out / "ranked_leads.xlsx")
    assert wb.sheetnames == ["Ranked", "Rejected", "Summary"]
    ranked = wb["Ranked"]
    assert ranked.freeze_panes == "A2"
    assert ranked["A1"].value == "Rank" and ranked["B1"].value == "Tier"
    assert ranked["B2"].fill.fgColor.rgb.endswith("C6EFCE")  # tier A color from icp.yaml
    assert wb["Rejected"].freeze_panes == "A2"


def test_output_is_deterministic(tmp_path, messy_leads):
    csv_path = tmp_path / "leads.csv"
    messy_leads.to_csv(csv_path, index=False)
    run(csv_path, ROOT / "icp.yaml", tmp_path / "a", run_date=RUN_DATE)
    run(csv_path, ROOT / "icp.yaml", tmp_path / "b", run_date=RUN_DATE)
    for name in ("ranked_leads.csv", "rejected_leads.csv", "summary.md"):
        assert (tmp_path / "a" / name).read_bytes() == (tmp_path / "b" / name).read_bytes()


def test_sample_data_runs(tmp_path):
    result = run(ROOT / "data" / "raw" / "sample_leads.csv", ROOT / "icp.yaml", tmp_path, run_date=RUN_DATE)
    s = result.stats
    assert s["rows_in"] == 500
    assert s["ranked"] + s["rejected"] + s["duplicates_removed"] == s["rows_in"]
    assert result.ranked["score"].is_monotonic_decreasing


def test_read_csv_errors(tmp_path):
    with pytest.raises(InputError, match="Can't find"):
        read_leads_csv(tmp_path / "missing.csv")
    empty = tmp_path / "empty.csv"
    empty.write_text("", encoding="utf-8")
    with pytest.raises(InputError, match="empty"):
        read_leads_csv(empty)
    header_only = tmp_path / "header.csv"
    header_only.write_text("company_name,email\n", encoding="utf-8")
    with pytest.raises(InputError, match="no leads"):
        read_leads_csv(header_only)


def test_read_csv_from_excel_ansi(tmp_path):
    path = tmp_path / "ansi.csv"
    path.write_bytes("company_name,contact_first_name,contact_last_name,email\nCafé Ltd,José,Núñez,j@cafe.example\n"
                     .encode("cp1252"))
    df = read_leads_csv(path)
    assert df.loc[0, "contact_first_name"] == "José"


def test_cli_success_and_friendly_errors(tmp_path, messy_leads, capsys):
    csv_path = tmp_path / "leads.csv"
    messy_leads.to_csv(csv_path, index=False)
    code = main(["run", "--input", str(csv_path), "--config", str(ROOT / "icp.yaml"), "--out", str(tmp_path / "o")])
    assert code == 0
    assert "Done!" in capsys.readouterr().out

    code = main(["run", "--input", str(tmp_path / "nope.csv"), "--config", str(ROOT / "icp.yaml")])
    assert code == 1
    assert "Problem: Can't find the lead file" in capsys.readouterr().err


def test_locked_excel_file_leaves_other_outputs_untouched(tmp_path, messy_leads, config, monkeypatch):
    import leadflow.pipeline as pipeline

    out = tmp_path / "o"
    out.mkdir()
    (out / "summary.md").write_text("previous run", encoding="utf-8")

    def locked(result, path):
        raise PermissionError(13, "Permission denied", str(path))

    monkeypatch.setattr(pipeline, "write_excel", locked)
    with pytest.raises(PermissionError):
        pipeline.write_outputs(process(messy_leads, config, run_date=RUN_DATE), out)
    assert (out / "summary.md").read_text(encoding="utf-8") == "previous run"
    assert not (out / "ranked_leads.csv").exists()


def test_ranked_csv_round_trips(tmp_path, messy_leads):
    csv_path = tmp_path / "leads.csv"
    messy_leads.to_csv(csv_path, index=False)
    run(csv_path, ROOT / "icp.yaml", tmp_path / "o", run_date=RUN_DATE)
    ranked = pd.read_csv(tmp_path / "o" / "ranked_leads.csv", encoding="utf-8-sig")
    assert ranked.columns[0] == "rank"
    assert (ranked["annual_revenue"].dropna() % 1 == 0).all()
